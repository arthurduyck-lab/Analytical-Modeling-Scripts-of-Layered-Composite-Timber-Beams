"""
Convert the legacy Excel input files (data_LAM.xlsx, data_CON.xlsx) into the
tab-separated text format now read by read_data().

Run it once per method folder, or from the repository root with --recursive.
Requires openpyxl only for the conversion itself; once converted, the analysis
scripts no longer depend on openpyxl.

Usage
-----
    python xlsx_to_csv.py                      # convert *.xlsx in current folder
    python xlsx_to_csv.py --recursive          # walk sub-folders too
    python xlsx_to_csv.py data_LAM.xlsx        # convert specific files
    python xlsx_to_csv.py --sheet Sheet1       # non-default sheet name
"""

import argparse
import glob
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required for the conversion: pip install openpyxl")


def convert(xlsx_path, sheet=None, sep='\t', overwrite=False):
    """Convert one .xlsx file to a separated-value text file."""
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]

    rows = [r for r in ws.iter_rows(values_only=True)
            if not all(v is None for v in r)]
    if not rows:
        raise ValueError(f"'{xlsx_path}' (sheet '{ws.title}') is empty.")

    # Drop trailing all-empty columns
    width = max(
        max((i + 1 for i, v in enumerate(r) if v is not None), default=0)
        for r in rows
    )
    rows = [r[:width] for r in rows]

    csv_path = os.path.splitext(xlsx_path)[0] + '.csv'
    if os.path.exists(csv_path) and not overwrite:
        raise FileExistsError(
            f"'{csv_path}' already exists. Use --overwrite to replace it."
        )

    with open(csv_path, 'w', encoding='utf-8', newline='') as f:
        for r in rows:
            cells = []
            for v in r:
                if v is None:
                    cells.append('')
                elif isinstance(v, float) and v.is_integer():
                    cells.append(str(int(v)))      # 30.0 -> 30
                else:
                    cells.append(str(v))
            f.write(sep.join(cells) + '\n')

    return csv_path, ws.title, len(rows) - 1


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('files', nargs='*',
                    help="Files to convert (default: every .xlsx in the folder)")
    ap.add_argument('--recursive', action='store_true',
                    help="Also convert .xlsx files found in sub-folders")
    ap.add_argument('--sheet', default=None,
                    help="Worksheet name (default: first sheet of the workbook)")
    ap.add_argument('--sep', default='\t',
                    help="Output separator (default: tabulation)")
    ap.add_argument('--overwrite', action='store_true',
                    help="Replace existing .csv files")
    args = ap.parse_args()

    if args.files:
        targets = args.files
    elif args.recursive:
        targets = glob.glob('**/*.xlsx', recursive=True)
    else:
        targets = glob.glob('*.xlsx')

    targets = [t for t in targets if not os.path.basename(t).startswith('~$')]

    if not targets:
        print("No .xlsx file found.")
        return

    n_ok = 0
    for xlsx in sorted(targets):
        try:
            csv_path, sheet, n_rows = convert(
                xlsx, sheet=args.sheet, sep=args.sep, overwrite=args.overwrite
            )
        except Exception as exc:
            print(f"  SKIPPED {xlsx}: {type(exc).__name__}: {exc}")
            continue
        print(f"  {xlsx}  ->  {csv_path}   (sheet '{sheet}', {n_rows} data row(s))")
        n_ok += 1

    print(f"\n{n_ok}/{len(targets)} file(s) converted.")


if __name__ == '__main__':
    main()
